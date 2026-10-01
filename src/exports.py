"""Excel workbook and HTML plan-summary exports.

Every export generated here is labelled "Synthetic educational example" and
contains only data entered for the synthetic demo households in this
simulator -- never real client data.
"""

from __future__ import annotations

import html
import io
import sqlite3
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src import calculations, models
from src.audit import get_events
from src.db import CHECKLIST_CATEGORIES

LABEL = "Synthetic educational example"

HEADER_FILL = PatternFill(start_color="1C5CAB", end_color="1C5CAB", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LABEL_FONT = Font(italic=True, color="B00000", bold=True)
TITLE_FONT = Font(bold=True, size=14)


def _write_label_and_title(ws: Worksheet, title: str) -> int:
    """Write the synthetic-data label and a sheet title. Returns next free row."""
    ws["A1"] = LABEL
    ws["A1"].font = LABEL_FONT
    ws["A2"] = title
    ws["A2"].font = TITLE_FONT
    return 4


def _write_table(ws: Worksheet, start_row: int, headers: list[str], rows: list[list]) -> int:
    """Write a header row (bold, filled) followed by data rows. Returns next free row."""
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left")

    for r_offset, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=start_row + r_offset, column=col_idx, value=value)

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24

    return start_row + len(rows) + 2


def _household_summary_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    household = models.get_household(conn, household_id)
    ws = wb.active
    ws.title = "Household Summary"
    row = _write_label_and_title(ws, f"Household Summary: {household['name']}")

    row = _write_table(
        ws,
        row,
        ["Field", "Value"],
        [
            ["Household name", household["name"]],
            ["Advisor", household["advisor"]],
            ["CRM stage", household["crm_stage"]],
            ["Risk tolerance", household["risk_tolerance"]],
            ["Contact preference", household["contact_preference"]],
            ["Next meeting date", household["next_meeting_date"] or "Not scheduled"],
            ["Created at", household["created_at"]],
            ["Notes", household["notes"] or ""],
        ],
    )

    people = models.list_people(conn, household_id)
    ws.cell(row=row, column=1, value="Household Members").font = Font(bold=True, size=12)
    row += 1
    _write_table(
        ws,
        row,
        ["Name", "Relationship", "Date of birth", "Dependant"],
        [[p["name"], p["relationship"], p["date_of_birth"] or "", "Yes" if p["is_dependant"] else "No"] for p in people],
    )


def _financial_snapshot_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    ws = wb.create_sheet("Financial Snapshot")
    row = _write_label_and_title(ws, "Financial Snapshot")

    s = models.get_financial_summary(conn, household_id)

    def pct(v):
        return "N/A" if v is None else f"{v * 100:.1f}%"

    def money(v):
        return "N/A" if v is None else round(v, 2)

    row = _write_table(
        ws,
        row,
        ["Metric", "Value"],
        [
            ["Total assets", money(s["total_assets"])],
            ["Total liabilities", money(s["total_liabilities"])],
            ["Net worth", money(s["net_worth"])],
            ["Liquid assets", money(s["liquid_assets"])],
            ["Monthly income", money(s["monthly_income"])],
            ["Monthly living expenses", money(s["monthly_expenses"])],
            ["Monthly debt payments", money(s["monthly_debt_payments"])],
            ["Monthly surplus", money(s["monthly_surplus"])],
            ["Savings rate", pct(s["savings_rate"])],
            ["Debt-to-income ratio", pct(s["debt_to_income"])],
            ["Emergency-fund coverage (months)", money(s["emergency_fund_months"])],
        ],
    )

    items = models.list_financial_items(conn, household_id)
    ws.cell(row=row, column=1, value="Financial Line Items").font = Font(bold=True, size=12)
    row += 1
    _write_table(
        ws,
        row,
        ["Category", "Subcategory", "Amount", "Frequency", "As of"],
        [[i["category"], i["subcategory"], i["amount"], i["frequency"], i["as_of_date"]] for i in items],
    )


def _goals_and_scenarios_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    ws = wb.create_sheet("Goals & Scenarios")
    row = _write_label_and_title(ws, "Goals and Scenario Projections")

    goals = models.list_goals(conn, household_id)
    goal_rows = []
    for g in goals:
        gap = calculations.goal_funding_gap(g["target_amount"], g["current_savings"])
        progress = calculations.goal_funding_progress(g["current_savings"], g["target_amount"])
        goal_rows.append(
            [
                g["name"],
                g["goal_type"],
                g["target_amount"],
                g["target_date"] or "",
                g["current_savings"],
                g["monthly_contribution"],
                gap,
                "N/A" if progress is None else f"{progress * 100:.1f}%",
            ]
        )
    row = _write_table(
        ws,
        row,
        ["Goal", "Type", "Target amount", "Target date", "Current savings", "Monthly contribution", "Funding gap", "Progress"],
        goal_rows,
    )

    ws.cell(row=row, column=1, value="Scenario Projections").font = Font(bold=True, size=12)
    row += 1
    scenario_rows = []
    for g in goals:
        assumptions = models.get_scenario_assumptions(conn, g["id"])
        if g["goal_type"] == "retirement":
            from datetime import date

            years = 20
            if g["target_date"]:
                try:
                    years = max((date.fromisoformat(g["target_date"]) - date.today()).days / 365, 0.5)
                except ValueError:
                    pass
            results = calculations.run_retirement_scenarios(
                g["current_savings"],
                g["monthly_contribution"],
                years,
                {k: calculations.ScenarioAssumption(k, v["annual_return"], v["withdrawal_rate"]) for k, v in assumptions.items()},
            )
            for scenario, r in results.items():
                scenario_rows.append(
                    [g["name"], scenario, r["annual_return"], r.get("withdrawal_rate"), round(r["projected_balance"], 2), round(r["annual_income"], 2)]
                )
    row = _write_table(
        ws,
        row,
        ["Goal", "Scenario", "Annual return", "Withdrawal rate", "Projected balance", "Est. annual income"],
        scenario_rows,
    )


def _checklists_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    ws = wb.create_sheet("Checklists")
    row = _write_label_and_title(ws, "Planning Checklists")
    items = models.list_checklist_items(conn, household_id)
    _write_table(
        ws,
        row,
        ["Category", "Item", "Status", "Note", "Updated at"],
        [[i["category"], i["item_name"], i["status"], i["note"] or "", i["updated_at"]] for i in items],
    )


def _meetings_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    ws = wb.create_sheet("Meetings")
    row = _write_label_and_title(ws, "Meetings")
    meetings = models.list_meetings(conn, household_id)
    _write_table(
        ws,
        row,
        ["Date", "Status", "Notes", "Recommendations"],
        [[m["meeting_date"], m["status"], m["notes"] or "", m["recommendations"] or ""] for m in meetings],
    )


def _tasks_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    ws = wb.create_sheet("Tasks")
    row = _write_label_and_title(ws, "Tasks")
    tasks = models.list_tasks(conn, household_id)
    _write_table(
        ws,
        row,
        ["Description", "Assigned to", "Due date", "Status", "Completed at"],
        [[t["description"], t["assigned_to"], t["due_date"] or "", t["status"], t["completed_at"] or ""] for t in tasks],
    )


def _audit_sheet(wb: Workbook, conn: sqlite3.Connection, household_id: int) -> None:
    ws = wb.create_sheet("Audit History")
    row = _write_label_and_title(ws, "Audit History")
    events = get_events(conn, household_id=household_id, limit=1000)
    _write_table(
        ws,
        row,
        ["Timestamp", "Actor", "Event type", "Entity type", "Description"],
        [[e["timestamp"], e["actor"], e["event_type"], e["entity_type"], e["description"]] for e in events],
    )


def generate_workbook(conn: sqlite3.Connection, household_id: int) -> bytes:
    """Build the advisor workbook for a household and return it as bytes."""
    wb = Workbook()
    _household_summary_sheet(wb, conn, household_id)
    _financial_snapshot_sheet(wb, conn, household_id)
    _goals_and_scenarios_sheet(wb, conn, household_id)
    _checklists_sheet(wb, conn, household_id)
    _meetings_sheet(wb, conn, household_id)
    _tasks_sheet(wb, conn, household_id)
    _audit_sheet(wb, conn, household_id)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _esc(value) -> str:
    return html.escape(str(value)) if value is not None else ""


def generate_plan_summary_html(conn: sqlite3.Connection, household_id: int) -> str:
    """Build a clean, printable client-facing plan summary as an HTML string."""
    household = models.get_household(conn, household_id)
    summary = models.get_financial_summary(conn, household_id)
    goals = models.list_goals(conn, household_id)
    checklist = models.list_checklist_items(conn, household_id)
    meetings = models.list_meetings(conn, household_id)
    latest_meeting = meetings[0] if meetings else None

    def money(v):
        return "N/A" if v is None else f"${v:,.2f}"

    def pct(v):
        return "N/A" if v is None else f"{v * 100:.1f}%"

    goal_rows = ""
    for g in goals:
        gap = calculations.goal_funding_gap(g["target_amount"], g["current_savings"])
        progress = calculations.goal_funding_progress(g["current_savings"], g["target_amount"])
        goal_rows += f"""
        <tr>
          <td>{_esc(g['name'])}</td>
          <td>{_esc(g['goal_type'].replace('_', ' ').title())}</td>
          <td>{money(g['target_amount'])}</td>
          <td>{_esc(g['target_date'] or '-')}</td>
          <td>{money(g['current_savings'])}</td>
          <td>{money(gap)}</td>
          <td>{pct(progress)}</td>
        </tr>"""

    checklist_rows = ""
    for category in CHECKLIST_CATEGORIES:
        cat_items = [c for c in checklist if c["category"] == category]
        if not cat_items:
            continue
        checklist_rows += f'<tr><td colspan="2" class="section-row"><strong>{_esc(category)}</strong></td></tr>'
        for item in cat_items:
            checklist_rows += f"""
            <tr>
              <td>{_esc(item['item_name'])}</td>
              <td>{_esc(item['status'])}</td>
            </tr>"""

    recommendations_html = ""
    if latest_meeting and latest_meeting["recommendations"]:
        recommendations_html = f"<p>{_esc(latest_meeting['recommendations'])}</p>"
    else:
        recommendations_html = "<p><em>No recommendations on file yet.</em></p>"

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="color-scheme" content="light">
<title>Plan Summary — {_esc(household['name'])}</title>
<style>
  html {{ color-scheme: light; background: #ffffff; }}
  body {{ font-family: Georgia, 'Times New Roman', serif; color: #1a1a1a; background: #ffffff;
         max-width: 800px; margin: 2rem auto; padding: 0 1.5rem; line-height: 1.5; }}
  .label {{ background: #fdecea; color: #a61b1b; border: 1px solid #a61b1b; padding: 0.5rem 1rem;
            border-radius: 4px; font-weight: bold; text-align: center; margin-bottom: 1.5rem; }}
  h1 {{ font-size: 1.8rem; margin-bottom: 0.2rem; }}
  h2 {{ font-size: 1.2rem; border-bottom: 2px solid #1c5cab; padding-bottom: 0.3rem; margin-top: 2rem; }}
  .meta {{ color: #555; font-size: 0.9rem; margin-bottom: 1.5rem; }}
  table {{ width: 100%; border-collapse: collapse; margin: 0.75rem 0; }}
  th, td {{ text-align: left; padding: 0.4rem 0.6rem; border-bottom: 1px solid #ddd; font-size: 0.95rem; }}
  th {{ background: #f0f4fa; }}
  .section-row td {{ background: #f7f7f5; padding-top: 0.6rem; }}
  .metric-grid {{ display: grid; grid-template-columns: repeat(2, 1fr); gap: 0.5rem 2rem; margin: 0.75rem 0; }}
  .metric {{ display: flex; justify-content: space-between; border-bottom: 1px dotted #ccc; padding: 0.2rem 0; }}
  .disclaimer {{ font-size: 0.85rem; color: #555; border-top: 1px solid #ccc; margin-top: 2rem; padding-top: 1rem; }}
  @media print {{ body {{ margin: 0; padding: 0 1rem; }} }}
</style>
</head>
<body>
  <div class="label">SYNTHETIC EDUCATIONAL EXAMPLE — NOT A REAL CLIENT OR FINANCIAL ADVICE</div>
  <h1>Financial Plan Summary</h1>
  <div class="meta">Prepared for <strong>{_esc(household['name'])}</strong> · Advisor: {_esc(household['advisor'])} · Generated {generated_at}</div>

  <h2>Financial Snapshot</h2>
  <div class="metric-grid">
    <div class="metric"><span>Total assets</span><strong>{money(summary['total_assets'])}</strong></div>
    <div class="metric"><span>Total liabilities</span><strong>{money(summary['total_liabilities'])}</strong></div>
    <div class="metric"><span>Net worth</span><strong>{money(summary['net_worth'])}</strong></div>
    <div class="metric"><span>Monthly income</span><strong>{money(summary['monthly_income'])}</strong></div>
    <div class="metric"><span>Monthly surplus</span><strong>{money(summary['monthly_surplus'])}</strong></div>
    <div class="metric"><span>Savings rate</span><strong>{pct(summary['savings_rate'])}</strong></div>
    <div class="metric"><span>Debt-to-income ratio</span><strong>{pct(summary['debt_to_income'])}</strong></div>
    <div class="metric"><span>Emergency-fund coverage</span><strong>{summary['emergency_fund_months'] and f"{summary['emergency_fund_months']:.1f} months" or "N/A"}</strong></div>
  </div>

  <h2>Goals &amp; Funding Progress</h2>
  <table>
    <tr><th>Goal</th><th>Type</th><th>Target</th><th>Target date</th><th>Current savings</th><th>Gap</th><th>Progress</th></tr>
    {goal_rows or '<tr><td colspan="7"><em>No goals recorded.</em></td></tr>'}
  </table>

  <h2>Planning Checklist Summary</h2>
  <table>
    <tr><th>Item</th><th>Status</th></tr>
    {checklist_rows or '<tr><td colspan="2"><em>No checklist items recorded.</em></td></tr>'}
  </table>

  <h2>Recommendations for Further Professional Review</h2>
  {recommendations_html}

  <div class="disclaimer">
    This document is a synthetic, educational example produced by a personal finance
    planning practice simulator. It does not represent a real client, account, or
    financial plan, and nothing in it is personalized financial, tax, legal, or
    insurance advice. Any recommendation above must be reviewed with a licensed
    professional before acting on it.
  </div>
</body>
</html>"""
